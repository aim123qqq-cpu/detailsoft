from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_ORIENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = r"C:\Users\USER\Documents\ChatGPT\DetailSoft\Памятка_цифровая_безопасность_подростков_в_стиле_примера.docx"
IMG = r"C:\Users\USER\Documents\ChatGPT\DetailSoft\digital_safety_illustration.png"

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr(); shd = tcPr.find(qn('w:shd'))
    if shd is None: shd = OxmlElement('w:shd'); tcPr.append(shd)
    shd.set(qn('w:fill'), fill)

def noborder(cell):
    tcPr = cell._tc.get_or_add_tcPr(); b = tcPr.first_child_found_in('w:tcBorders')
    if b is None: b = OxmlElement('w:tcBorders'); tcPr.append(b)
    for e in ('top','left','bottom','right','insideH','insideV'):
        x = OxmlElement('w:'+e); x.set(qn('w:val'),'nil'); b.append(x)

def margins(cell, top=50, start=80, bottom=50, end=80):
    tcPr=cell._tc.get_or_add_tcPr(); m=tcPr.first_child_found_in('w:tcMar')
    if m is None: m=OxmlElement('w:tcMar'); tcPr.append(m)
    for k,v in [('top',top),('start',start),('bottom',bottom),('end',end)]:
        n=OxmlElement('w:'+k); n.set(qn('w:w'),str(v)); n.set(qn('w:type'),'dxa'); m.append(n)

def set_font(run,size=11,bold=False,italic=False,color='171717'):
    run.font.name='Times New Roman'; run._element.rPr.rFonts.set(qn('w:ascii'),'Times New Roman'); run._element.rPr.rFonts.set(qn('w:hAnsi'),'Times New Roman')
    run.font.size=Pt(size); run.bold=bold; run.italic=italic; run.font.color.rgb=RGBColor.from_string(color)

def add_para(cell,text='',size=10.4,bold=False,italic=False,align=None,after=2,first=0):
    p=cell.add_paragraph() if cell.paragraphs[0].text else cell.paragraphs[0]
    p.paragraph_format.space_after=Pt(after); p.paragraph_format.line_spacing=1.0; p.paragraph_format.first_line_indent=Inches(first)
    if align is not None: p.alignment=align
    r=p.add_run(text); set_font(r,size,bold,italic); return p

def add_bullet(cell,text):
    p=cell.add_paragraph(); p.paragraph_format.space_after=Pt(2); p.paragraph_format.line_spacing=1.0; p.paragraph_format.left_indent=Inches(.12); p.paragraph_format.first_line_indent=Inches(-.12)
    r=p.add_run('➢ '); set_font(r,10.5,False,False,'222222'); r=p.add_run(text); set_font(r,10.2); return p

doc=Document(); sec=doc.sections[0]; sec.orientation=WD_ORIENT.LANDSCAPE; sec.page_width=Inches(11.69); sec.page_height=Inches(8.27); sec.top_margin=Inches(.42); sec.bottom_margin=Inches(.42); sec.left_margin=Inches(.45); sec.right_margin=Inches(.45)
style=doc.styles['Normal']; style.font.name='Times New Roman'; style._element.rPr.rFonts.set(qn('w:ascii'),'Times New Roman'); style._element.rPr.rFonts.set(qn('w:hAnsi'),'Times New Roman'); style.font.size=Pt(10.2)

p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY; p.paragraph_format.space_after=Pt(5); p.paragraph_format.line_spacing=1.0
r=p.add_run('Интернет может быть прекрасным и полезным средством для обучения, отдыха или общения с друзьями. Но он может быть опасен: мошенники, кибербуллинг, вредоносные программы и попытки выманить личные данные встречаются и в социальных сетях, и в играх. Важно знать основные правила цифровой безопасности и понимать, к кому обратиться, если что-то пошло не так.'); set_font(r,11)

bar=doc.add_table(rows=1,cols=3); bar.alignment=WD_TABLE_ALIGNMENT.CENTER; bar.autofit=False
for c,w,col in zip(bar.rows[0].cells,[3.0,3.0,4.8],['E7B400','F07818','16264A']):
    c.width=Inches(w); shade(c,col); noborder(c); margins(c,45,0,45,0)

main=doc.add_table(rows=1,cols=3); main.alignment=WD_TABLE_ALIGNMENT.CENTER; main.autofit=False
for c,w in zip(main.rows[0].cells,[3.0,4.15,3.65]): c.width=Inches(w); noborder(c); margins(c,80,100,40,100); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.TOP

left=main.cell(0,0); left.paragraphs[0].alignment=WD_ALIGN_PARAGRAPH.CENTER; left.paragraphs[0].add_run().add_picture(IMG,width=Inches(2.35))
add_para(left,'Почему важно говорить о безопасности',10.4,True,False,WD_ALIGN_PARAGRAPH.CENTER,4)
add_para(left,'Подростки проводят в сети много времени: учатся, общаются, смотрят видео и играют. Поэтому полезно заранее договориться о правилах и знать признаки опасной ситуации.',10.1,False,True,WD_ALIGN_PARAGRAPH.JUSTIFY,3)
for t in ['Вам пишут незнакомые люди и просят прислать фото или код из СМС.','Кто-то обещает подарок, выигрыш или лёгкий заработок за личные данные.','После общения в сети вы чувствуете страх, стыд или давление.','Вас торопят перейти по ссылке, установить программу или встретиться лично.']:
    add_bullet(left,t)

mid=main.cell(0,1); add_para(mid,'Правила безопасного поведения в Интернете',12,True,False,WD_ALIGN_PARAGRAPH.CENTER,5)
add_para(mid,'1. Защищайте личную информацию',10.8,True,False,None,1)
for t in ['Не публикуйте адрес, номер телефона, пароль, коды из СМС, документы и точное местоположение.','Используйте разные длинные пароли и включите двухэтапную проверку входа.','Проверьте настройки приватности: личные публикации должны быть доступны только тем, кому вы доверяете.']:
    add_bullet(mid,t)
add_para(mid,'2. Проверяйте сообщения и ссылки',10.8,True,False,None,1)
for t in ['Не открывайте подозрительные ссылки и файлы, даже если они пришли от знакомого аккаунта.','Сверяйте адрес сайта и не вводите пароль на странице, в которой сомневаетесь.','Если сообщение пугает или торопит, остановитесь и покажите его взрослому.']:
    add_bullet(mid,t)
add_para(mid,'3. Общайтесь уважительно',10.8,True,False,None,1)
for t in ['Не отвечайте агрессией на агрессию и не пересылайте обидные публикации.','При травле сохраните скриншоты, заблокируйте обидчика и расскажите взрослому.']:
    add_bullet(mid,t)

right=main.cell(0,2); add_para(right,'Если возникла проблема',12,True,False,WD_ALIGN_PARAGRAPH.CENTER,5)
add_para(right,'Сохраняйте спокойствие и действуйте по шагам:',10.5,False,True,None,3)
for t in ['не продолжайте переписку с человеком, который давит или угрожает;','сделайте скриншоты и сохраните ссылку на страницу;','заблокируйте пользователя и пожалуйтесь на контент;','обратитесь к родителю, педагогу или другому взрослому, которому доверяете;','если есть угроза жизни или здоровью — звоните 112.']:
    add_bullet(right,t)
add_para(right,'Памятка для подростка',10.8,True,False,WD_ALIGN_PARAGRAPH.RIGHT,2)
add_para(right,'Безопасность начинается с простого правила: думай перед тем, как нажать, отправить или встретиться.',10.4,False,True,WD_ALIGN_PARAGRAPH.JUSTIFY,4)
right.paragraphs[-1].alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
right.paragraphs[0].alignment=WD_ALIGN_PARAGRAPH.CENTER

bar2=doc.add_table(rows=1,cols=3); bar2.alignment=WD_TABLE_ALIGNMENT.CENTER; bar2.autofit=False
for c,w,col in zip(bar2.rows[0].cells,[3.0,3.0,4.8],['E7B400','F07818','16264A']): c.width=Inches(w); shade(c,col); noborder(c); margins(c,35,0,35,0)
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.RIGHT; p.paragraph_format.space_before=Pt(3); p.paragraph_format.space_after=Pt(0); r=p.add_run('Цифровая безопасность: полезные привычки на каждый день'); set_font(r,9.2,False,True,'444444')
doc.save(OUT); print(OUT)
