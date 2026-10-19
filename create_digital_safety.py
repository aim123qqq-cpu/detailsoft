from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = r"C:\Users\USER\Documents\ChatGPT\DetailSoft\Памятка_цифровая_безопасность_подростков.docx"
IMG = r"C:\Users\USER\Documents\ChatGPT\DetailSoft\digital_safety_illustration.png"

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr(); shd = tcPr.find(qn('w:shd'))
    if shd is None: shd = OxmlElement('w:shd'); tcPr.append(shd)
    shd.set(qn('w:fill'), fill)

def border(cell, color='D9D9D9', sz='8'):
    tcPr = cell._tc.get_or_add_tcPr(); borders = tcPr.first_child_found_in('w:tcBorders')
    if borders is None: borders = OxmlElement('w:tcBorders'); tcPr.append(borders)
    for edge in ('top','left','bottom','right','insideH','insideV'):
        el = borders.find(qn('w:'+edge))
        if el is None: el = OxmlElement('w:'+edge); borders.append(el)
        el.set(qn('w:val'),'single'); el.set(qn('w:sz'),sz); el.set(qn('w:space'),'0'); el.set(qn('w:color'),color)

def margins(cell, top=100, start=140, bottom=100, end=140):
    tcPr = cell._tc.get_or_add_tcPr(); tcMar = tcPr.first_child_found_in('w:tcMar')
    if tcMar is None: tcMar = OxmlElement('w:tcMar'); tcPr.append(tcMar)
    for m,v in [('top',top),('start',start),('bottom',bottom),('end',end)]:
        node = tcMar.find(qn('w:'+m))
        if node is None: node = OxmlElement('w:'+m); tcMar.append(node)
        node.set(qn('w:w'),str(v)); node.set(qn('w:type'),'dxa')

def font(run, size=11, color='1C2340', bold=False, italic=False):
    run.font.name='Aptos'; run._element.rPr.rFonts.set(qn('w:ascii'),'Aptos'); run._element.rPr.rFonts.set(qn('w:hAnsi'),'Aptos')
    run.font.size=Pt(size); run.font.color.rgb=RGBColor.from_string(color); run.bold=bold; run.italic=italic

doc=Document(); sec=doc.sections[0]
sec.page_width=Inches(8.27); sec.page_height=Inches(11.69); sec.top_margin=Inches(.38); sec.bottom_margin=Inches(.38); sec.left_margin=Inches(.48); sec.right_margin=Inches(.48)
normal=doc.styles['Normal']; normal.font.name='Aptos'; normal._element.rPr.rFonts.set(qn('w:ascii'),'Aptos'); normal._element.rPr.rFonts.set(qn('w:hAnsi'),'Aptos'); normal.font.size=Pt(10)

top=doc.add_table(rows=1,cols=2); top.alignment=WD_TABLE_ALIGNMENT.CENTER; top.autofit=False
top.columns[0].width=Inches(5.3); top.columns[1].width=Inches(2.0)
for c in top.rows[0].cells: shade(c,'F2F4FF'); border(c,'F2F4FF','0'); margins(c,150,180,150,180); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
p=top.cell(0,0).paragraphs[0]; p.paragraph_format.space_after=Pt(2); r=p.add_run('ЦИФРОВАЯ БЕЗОПАСНОСТЬ'); font(r,24,'142B73',True)
p=top.cell(0,0).add_paragraph(); p.paragraph_format.space_after=Pt(5); r=p.add_run('Памятка для подростка'); font(r,14,'F05A47',True)
p=top.cell(0,0).add_paragraph(); p.paragraph_format.space_after=Pt(0); r=p.add_run('Интернет становится безопаснее, когда ты знаешь несколько простых правил.'); font(r,10.5,'3D4765')
top.cell(0,1).paragraphs[0].alignment=WD_ALIGN_PARAGRAPH.CENTER; top.cell(0,1).paragraphs[0].add_run().add_picture(IMG,width=Inches(1.7))

sp=doc.add_paragraph(); sp.paragraph_format.space_after=Pt(2)
intro=doc.add_paragraph(); intro.paragraph_format.space_after=Pt(7); intro.paragraph_format.left_indent=Inches(.08)
r=intro.add_run('Главное: '); font(r,11.5,'142B73',True); r=intro.add_run('не делись тем, что может навредить тебе, и не бойся просить помощи у взрослого.'); font(r,11.5)

rules=[('1','Проверь ссылку','Не открывай подозрительные ссылки и файлы, даже если их прислал знакомый аккаунт.','E8F7FF','00A6C7'),('2','Спрячь личное','Пароль, адрес, номер телефона и коды из СМС — только для тебя и родителей.','FFF4D6','F3A712'),('3','Включи защиту','Сделай длинный уникальный пароль и включи двухэтапную проверку входа.','F0EAFE','8B5CF6'),('4','Думай перед публикацией','Фото и сообщения могут сохранить и переслать. Спроси себя: «Мне будет спокойно, если это увидят другие?»','EAF9F0','2EAD68'),('5','Расскажи о проблеме','Если тебя пугают, шантажируют или взломали аккаунт — сохрани скриншоты и обратись к взрослому.','FFECEC','F05A47')]
tbl=doc.add_table(rows=3,cols=2); tbl.alignment=WD_TABLE_ALIGNMENT.CENTER; tbl.autofit=False
for row in tbl.rows:
    row.cells[0].width=Inches(3.55); row.cells[1].width=Inches(3.55)
idx=0
for i in range(3):
    for j in range(2):
        c=tbl.cell(i,j); border(c); margins(c,100,130,100,130); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.TOP
        if idx < len(rules):
            n,head,body,fill,accent=rules[idx]; shade(c,fill)
            p=c.paragraphs[0]; p.paragraph_format.space_after=Pt(2); r=p.add_run(n+'  '); font(r,17,accent,True); r=p.add_run(head); font(r,12,'1C2340',True)
            p=c.add_paragraph(); p.paragraph_format.space_after=Pt(0); p.paragraph_format.line_spacing=1.0; r=p.add_run(body); font(r,9.6,'36415F')
        else: shade(c,'FFFFFF'); c.text=''
        idx+=1

sp=doc.add_paragraph(); sp.paragraph_format.space_after=Pt(2)
check=doc.add_table(rows=1,cols=1); check.alignment=WD_TABLE_ALIGNMENT.CENTER
c=check.cell(0,0); shade(c,'142B73'); border(c,'142B73','0'); margins(c,120,180,120,180)
p=c.paragraphs[0]; p.paragraph_format.space_after=Pt(3); r=p.add_run('МОЙ БЫСТРЫЙ ЧЕК‑ЛИСТ'); font(r,11.5,'FFFFFF',True)
p=c.add_paragraph(); p.paragraph_format.space_after=Pt(0); p.paragraph_format.line_spacing=1.0; r=p.add_run('□  Пароль не повторяется на разных сайтах    □  Профиль закрыт для незнакомых    □  Я знаю, к кому обратиться за помощью'); font(r,10,'FFFFFF')
foot=doc.add_paragraph(); foot.alignment=WD_ALIGN_PARAGRAPH.CENTER; foot.paragraph_format.space_before=Pt(5); foot.paragraph_format.space_after=Pt(0); r=foot.add_run('Будь внимателен онлайн — и оставляй место для жизни офлайн.'); font(r,9.2,'6D7694',False,True)
doc.save(OUT); print(OUT)
